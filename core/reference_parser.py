import os
import logging
from typing import Dict, Any, List, Optional
from pathlib import Path
import docx
from docx.shared import Inches, Pt
import pptx

logger = logging.getLogger(__name__)

class ReferenceParser:
    """
    Parses uploaded DOCX or PPTX reference files to extract structure,
    tone, styling parameters, and headings so new works can be generated
    in the same likeness.
    """

    @staticmethod
    def parse_file(file_path: Path) -> Dict[str, Any]:
        suffix = file_path.suffix.lower()
        if suffix == ".docx":
            return ReferenceParser.parse_docx(file_path)
        elif suffix == ".pptx":
            return ReferenceParser.parse_pptx(file_path)
        else:
            raise ValueError(f"Неподдерживаемый формат образца: {suffix}. Ожидается .docx или .pptx")

    @staticmethod
    def parse_docx(file_path: Path) -> Dict[str, Any]:
        doc = docx.Document(file_path)
        headings = []
        sample_paragraphs = []
        fonts_detected = set()
        sizes_detected = set()

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Check heading style or uppercase formatting
            style_name = para.style.name.lower()
            if "heading" in style_name or "заголовок" in style_name or (len(text) < 120 and text.isupper()):
                headings.append(text)
            elif len(sample_paragraphs) < 15 and len(text) > 40:
                sample_paragraphs.append(text)

            for run in para.runs:
                if run.font.name:
                    fonts_detected.add(run.font.name)
                if run.font.size:
                    sizes_detected.add(run.font.size.pt)

        # Margins extraction
        margins = {}
        if doc.sections:
            sec = doc.sections[0]
            margins = {
                "top_mm": round(sec.top_margin.mm, 1) if sec.top_margin else 20,
                "bottom_mm": round(sec.bottom_margin.mm, 1) if sec.bottom_margin else 20,
                "left_mm": round(sec.left_margin.mm, 1) if sec.left_margin else 30,
                "right_mm": round(sec.right_margin.mm, 1) if sec.right_margin else 10
            }

        sample_text = "\n\n".join(sample_paragraphs[:8])
        font_name = list(fonts_detected)[0] if fonts_detected else "Times New Roman"

        return {
            "type": "docx",
            "file_name": file_path.name,
            "headings": headings[:15],
            "sample_text": sample_text,
            "font_name": font_name,
            "margins": margins,
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables)
        }

    @staticmethod
    def parse_pptx(file_path: Path) -> Dict[str, Any]:
        prs = pptx.Presentation(file_path)
        slides_info = []

        for idx, slide in enumerate(prs.slides, 1):
            title = ""
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    tf_text = shape.text_frame.text.strip()
                    if not tf_text:
                        continue
                    if shape == slide.shapes.title or not title:
                        title = tf_text
                    else:
                        texts.append(tf_text)
            slides_info.append({
                "slide_number": idx,
                "title": title or f"Слайд {idx}",
                "content_preview": texts[:3]
            })

        return {
            "type": "pptx",
            "file_name": file_path.name,
            "slide_count": len(prs.slides),
            "slides": slides_info[:12],
            "sample_text": "\n".join([f"{s['title']}: {', '.join(s['content_preview'])}" for s in slides_info[:6]])
        }
